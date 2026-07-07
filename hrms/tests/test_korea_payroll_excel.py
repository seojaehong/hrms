"""급여대장 엑셀 파싱·무결성 검증 코어 테스트 — framework-free.

실행: python3 hrms/tests/test_korea_payroll_excel.py
인메모리 xlsx(openpyxl)로 파싱·무결성(earnings합=세전, 세전-공제=실지급)을 검증한다.
"""

from __future__ import annotations

import importlib.util
import io
import pathlib
import unittest

import openpyxl

ROOT = pathlib.Path(__file__).resolve().parents[2]
MODULE_PATH = ROOT / "hrms" / "regional" / "south_korea" / "payroll_excel.py"


def load_module():
    spec = importlib.util.spec_from_file_location("korea_payroll_excel", MODULE_PATH)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


mod = load_module()
PROFILE = mod.PROFILES["monthly"]


def build_workbook(employees: list[dict], sheet: str | None = None) -> bytes:
    """monthly 프로파일 컬럼 매핑에 맞춰 인메모리 xlsx 바이트를 만든다.

    각 employee: {name, dept, join, email, earnings{label:amount}, deductions{label:amount},
                  gross, ded_total, net}. 지정된 셀만 채우고 나머지는 비운다.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet or PROFILE["sheet"]
    cols = PROFILE["cols"]
    earn_by_label = {label: c for c, label in PROFILE["earnings"].items()}
    ded_by_label = {label: c for c, label in PROFILE["deductions"].items()}

    ws.cell(1, cols["name"]).value = "성명"  # 헤더 행 (row 1)은 파싱에서 스킵됨
    row = 2
    for e in employees:
        ws.cell(row, cols["name"]).value = e["name"]
        ws.cell(row, cols["dept"]).value = e.get("dept", "")
        ws.cell(row, cols["join"]).value = e.get("join", "")
        ws.cell(row, cols["email"]).value = e.get("email", "")
        for label, amount in e.get("earnings", {}).items():
            ws.cell(row, earn_by_label[label]).value = amount
        for label, amount in e.get("deductions", {}).items():
            ws.cell(row, ded_by_label[label]).value = amount
        ws.cell(row, cols["gross"]).value = e.get("gross", 0)
        ws.cell(row, cols["ded_total"]).value = e.get("ded_total", 0)
        ws.cell(row, cols["net"]).value = e.get("net", 0)
        row += 1

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


class TestParse(unittest.TestCase):
    def test_parses_from_bytes_and_maps_columns(self):
        data = build_workbook([
            {"name": "김월급", "dept": "관리", "join": "2025-03-01", "email": "kim@x.com",
             "earnings": {"기본급": 3000000, "식대(비과세)": 200000},
             "deductions": {"소득세": 50000, "국민연금": 135000},
             "gross": 3200000, "ded_total": 185000, "net": 3015000},
        ])
        rows = mod.parse_payroll_workbook(data, "monthly")
        self.assertEqual(len(rows), 1)
        r = rows[0]
        self.assertEqual(r["name"], "김월급")
        self.assertEqual(r["dept"], "관리")
        self.assertEqual(r["join"], "2025-03-01")
        self.assertEqual(r["email"], "kim@x.com")
        self.assertEqual(r["earnings"], {"기본급": 3000000, "식대(비과세)": 200000})
        self.assertEqual(r["expected_gross"], 3200000)
        self.assertEqual(r["expected_net"], 3015000)

    def test_skips_summary_row(self):
        data = build_workbook([
            {"name": "이직원", "earnings": {"기본급": 2000000}, "deductions": {"소득세": 10000},
             "gross": 2000000, "ded_total": 10000, "net": 1990000},
            {"name": "합계", "earnings": {"기본급": 2000000}, "deductions": {"소득세": 10000},
             "gross": 2000000, "ded_total": 10000, "net": 1990000},
        ])
        rows = mod.parse_payroll_workbook(data, "monthly")
        self.assertEqual([r["name"] for r in rows], ["이직원"])

    def test_unknown_sheet_type_raises(self):
        with self.assertRaises(ValueError):
            mod.parse_payroll_workbook(b"", "weekly")


class TestValidateAndExtract(unittest.TestCase):
    def _rows(self, employees):
        return mod.parse_payroll_workbook(build_workbook(employees), "monthly")

    def test_valid_rows_pass_and_extract_returns_payload(self):
        employees = [
            {"name": "A", "earnings": {"기본급": 3000000}, "deductions": {"소득세": 100000},
             "gross": 3000000, "ded_total": 100000, "net": 2900000},
            {"name": "B", "earnings": {"기본급": 2500000, "보안수당": 100000},
             "deductions": {"소득세": 50000, "고용보험": 20000},
             "gross": 2600000, "ded_total": 70000, "net": 2530000},
            {"name": "C", "earnings": {"기본급": 1000000}, "deductions": {},
             "gross": 1000000, "ded_total": 0, "net": 1000000},
        ]
        rows = self._rows(employees)
        self.assertEqual(mod.validate_payroll_rows(rows), [])
        payload = mod.extract_payroll(build_workbook(employees), "2026-05", "monthly")
        self.assertEqual(payload["period"], "2026-05")
        self.assertEqual(payload["count"], 3)

    def test_earnings_sum_mismatch_flagged(self):
        rows = self._rows([
            {"name": "불일치", "earnings": {"기본급": 3000000}, "deductions": {"소득세": 100000},
             "gross": 3100000, "ded_total": 100000, "net": 3000000},  # earnings 합 3,000,000 != 세전 3,100,000
        ])
        errors = mod.validate_payroll_rows(rows)
        self.assertEqual(len(errors), 1)
        self.assertIn("earnings 합", errors[0])
        with self.assertRaises(ValueError):
            mod.extract_payroll(build_workbook([
                {"name": "불일치", "earnings": {"기본급": 3000000}, "deductions": {"소득세": 100000},
                 "gross": 3100000, "ded_total": 100000, "net": 3000000},
            ]), "2026-05", "monthly")

    def test_net_mismatch_flagged(self):
        rows = self._rows([
            {"name": "실지급오류", "earnings": {"기본급": 3000000}, "deductions": {"소득세": 100000},
             "gross": 3000000, "ded_total": 100000, "net": 2999999},  # 세전-공제=2,900,000 != 실지급
        ])
        errors = mod.validate_payroll_rows(rows)
        self.assertEqual(len(errors), 1)
        self.assertIn("세전-공제", errors[0])


if __name__ == "__main__":
    unittest.main(verbosity=2)
