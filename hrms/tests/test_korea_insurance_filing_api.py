"""4대보험 신고서 Frappe 연동 API 테스트 — framework-free.

실행: python3 hrms/tests/test_korea_insurance_filing_api.py

- human_approved=False → blocked (파일 미생성, frappe 없이 로드 가능)
- human_approved=True + fake frappe 스텁(직원 픽스처) → 취득 신고서 생성 및 셀 검증
템플릿 원본은 repo 밖 경로 참조(커밋 금지). 산출물은 tempfile(자동 삭제).
"""

from __future__ import annotations

import copy
import importlib.util
import pathlib
import sys
import tempfile
import unittest

import openpyxl

MODULE_PATH = (
	pathlib.Path(__file__).resolve().parents[1]
	/ "regional"
	/ "south_korea"
	/ "insurance_filing_api.py"
)

TEMPLATE_DIR = pathlib.Path(
	"C:/Users/iceam/OneDrive/_10_고객/_active/급여자동화/templates"
)
ACQUISITION_TEMPLATE = TEMPLATE_DIR / "근로자고용취득신고_전자신고용 (3).xlsx"

RRN_FIELD = "custom_resident_registration_number"


class FakeFrappe:
	"""get_all 만 제공하는 최소 스텁 (frappe.whitelist no-op 포함)."""

	def __init__(self, employees):
		self.employees = [copy.deepcopy(e) for e in employees]
		self.get_all_calls = []

	def whitelist(self):
		def decorator(fn):
			return fn

		return decorator

	def get_all(self, doctype, *, filters=None, fields=None):
		self.get_all_calls.append({"doctype": doctype, "filters": copy.deepcopy(filters), "fields": list(fields or [])})
		if doctype == "Employee":
			return [{f: copy.deepcopy(row.get(f)) for f in fields or []} for row in self.employees]
		return []


def load_module(fake_frappe=None):
	old_frappe = sys.modules.get("frappe")
	if fake_frappe is not None:
		sys.modules["frappe"] = fake_frappe
	else:
		sys.modules.pop("frappe", None)
	try:
		spec = importlib.util.spec_from_file_location("korea_insurance_filing_api", MODULE_PATH)
		module = importlib.util.module_from_spec(spec)
		assert spec.loader is not None
		spec.loader.exec_module(module)
		return module
	finally:
		if old_frappe is not None:
			sys.modules["frappe"] = old_frappe
		else:
			sys.modules.pop("frappe", None)


class TestBlockedGate(unittest.TestCase):
	def test_blocked_without_approval_and_no_file(self):
		# frappe 없이 로드 가능해야 하고, 승인 없으면 파일도 안 만든다.
		mod = load_module(fake_frappe=None)
		with tempfile.TemporaryDirectory() as tmp:
			result = mod.generate_insurance_filing(
				"acquisition",
				2026,
				7,
				str(ACQUISITION_TEMPLATE),
				human_approved=False,
				out_dir=tmp,
			)
			self.assertEqual(result["status"], "blocked")
			self.assertTrue(result["requires_human_confirmation"])
			# 산출물 디렉터리에 어떤 파일도 생성되지 않음
			self.assertEqual(list(pathlib.Path(tmp).iterdir()), [])

	def test_string_false_is_blocked(self):
		mod = load_module(fake_frappe=None)
		with tempfile.TemporaryDirectory() as tmp:
			result = mod.generate_insurance_filing(
				"acquisition", 2026, 7, str(ACQUISITION_TEMPLATE), human_approved="false", out_dir=tmp
			)
			self.assertEqual(result["status"], "blocked")


@unittest.skipUnless(ACQUISITION_TEMPLATE.exists(), f"템플릿 없음: {ACQUISITION_TEMPLATE}")
class TestApprovedAcquisition(unittest.TestCase):
	def test_creates_acquisition_report_from_site_employees(self):
		employees = [
			{
				"name": "HR-EMP-001",
				"employee_name": "김취득",
				"date_of_joining": "2026-07-15",
				"relieving_date": None,
				"employment_type": "정규직",
				RRN_FIELD: "0000009999999",
			},
			{
				"name": "HR-EMP-002",
				"employee_name": "박입사",
				"date_of_joining": "2026-07-01",
				"relieving_date": None,
				"employment_type": "정규직",
				# rrn 없음 → rrn_missing
			},
			{
				# 타 월 입사 → 대상 제외
				"name": "HR-EMP-003",
				"employee_name": "이전달",
				"date_of_joining": "2026-06-10",
				"relieving_date": None,
				"employment_type": "정규직",
			},
		]
		fake = FakeFrappe(employees)
		mod = load_module(fake_frappe=fake)
		with tempfile.TemporaryDirectory() as tmp:
			result = mod.generate_insurance_filing(
				"acquisition",
				2026,
				7,
				str(ACQUISITION_TEMPLATE),
				human_approved=True,
				out_dir=tmp,
				rrn_field=RRN_FIELD,
			)
			self.assertEqual(result["status"], "created")
			self.assertEqual(result["candidate_count"], 2)  # 7월 입사자 2명(타 월 제외)
			self.assertEqual(result["rrn_missing"], ["박입사"])
			self.assertTrue(pathlib.Path(result["file_path"]).exists())

			wb = openpyxl.load_workbook(result["file_path"])
			ws = wb["서식"]
			# detect_acquisitions 정렬: 취득일 오름차순 → 박입사(07-01), 김취득(07-15)
			self.assertEqual(ws["B3"].value, "박입사")
			self.assertEqual(ws["B4"].value, "김취득")
			self.assertEqual(ws["H3"].value, "20260701")
			self.assertEqual(ws["H4"].value, "20260715")
			# rrn 병합: 김취득만 존재
			self.assertEqual(ws["A4"].value, "0000009999999")
			self.assertIn(ws["A3"].value, (None, ""))
			wb.close()
		# Employee 조회가 실제로 일어났는지 확인
		self.assertEqual(fake.get_all_calls[0]["doctype"], "Employee")


if __name__ == "__main__":
	unittest.main()
