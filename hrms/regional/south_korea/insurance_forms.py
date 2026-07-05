"""4대보험 신고서 xlsx 생성기 — framework-free.

US-001 실측 매핑(docs/korea_hrms/insurance-templates-map.md)을 단일 출처로 사용해,
템플릿 원본을 복사한 뒤 데이터 행을 채워 실파일(xlsx)을 생성한다.

안전 불변식 (스킬 사고이력 2026-05-15):
  - candidates가 비어 있으면 ValueError — 빈 신고서를 만들지 않는다(fail-closed).
  - 사람 확인 게이트(human_approved)는 상위 _api 래퍼가 책임진다. 이 코어는 순수 생성만 한다.
  - 주민번호(rrn)는 후보에 있을 때만 채우고, 없으면 빈칸으로 둔다.
"""

from __future__ import annotations

import calendar
import datetime as dt
import importlib.util
import pathlib
from typing import Any

import openpyxl
from openpyxl.utils import get_column_letter

# 취득신고서 '서식' 시트 매핑 (US-001)
ACQUISITION_SHEET = "서식"
ACQUISITION_DATA_START_ROW = 3
# 취득일은 4개 보험(H/P/V/AD)에 동일, 보수월액은 G/O/U/AC에 동일 기입
ACQUISITION_DATE_COLS = ("H", "P", "V", "AD")
ACQUISITION_WAGE_COLS = ("G", "O", "U", "AC")

# 상실신고서 '서식' 시트 매핑 (US-001)
LOSS_SHEET = "서식"
LOSS_DATA_START_ROW = 2
# 상실일은 4개 보험(F=국민연금/I=건강/O=고용/T=산재)에 동일 기입
LOSS_DATE_COLS = ("F", "I", "O", "T")
LOSS_REASON_COL = "P"  # 고용보험상실사유구분코드

# 근로내용확인(일용직) '서식' 시트 매핑 (US-001)
DAILY_SHEET = "서식"
DAILY_DATA_START_ROW = 2
DAILY_DATA_LAST_ROW = 8  # 템플릿 잔존 데이터 행(2~8) — 쓰기 전 반드시 비운다(PII)
DAILY_CLEAR_LAST_COL = 46  # A~AT
DAILY_INSURANCE_COL = "A"  # 보험구분(기본 5)
DAILY_NAME_COL = "B"
DAILY_RRN_COL = "C"
DAILY_JOB_COL = "I"  # 직종코드
DAILY_WORKDAYS_COL = "AO"  # 근로일수 =AQ
DAILY_AVG_HOURS_COL = "AP"  # 일평균근로시간
DAILY_BASE_DAYS_COL = "AQ"  # 보수지급기초일수 =COUNT(J:AN)
DAILY_PAY_TOTAL_COL = "AR"  # 보수총액(과세소득)
DAILY_WAGE_TOTAL_COL = "AS"  # 임금총액 =AR
DAILY_SEPARATION_COL = "AT"  # 이직사유코드
DAILY_FIRST_DAY_COL_INDEX = 9  # 일자 d → 열 인덱스 9+d (J=10=1일)


def _load_insurance_filing():
	"""사이드카 코어(insurance_filing.py)를 패키지 import 없이 직접 로드."""
	path = pathlib.Path(__file__).with_name("insurance_filing.py")
	spec = importlib.util.spec_from_file_location("korea_insurance_filing", path)
	module = importlib.util.module_from_spec(spec)
	spec.loader.exec_module(module)
	return module


def _yyyymmdd(iso_value: Any) -> str:
	"""ISO date(YYYY-MM-DD) 또는 date → 신고 서식용 'YYYYMMDD' 문자열."""
	if iso_value in (None, ""):
		return ""
	if isinstance(iso_value, dt.date):
		return iso_value.strftime("%Y%m%d")
	day = dt.date.fromisoformat(str(iso_value)[:10])
	return day.strftime("%Y%m%d")


def generate_acquisition_report(
	template_path: str,
	candidates: list[dict],
	out_path: str,
	workplace_info: dict | None = None,
) -> str:
	"""취득신고서 xlsx 생성.

	candidates: insurance_filing.detect_acquisitions 출력
	  [{employee, employee_name, acquisition_date(ISO), monthly_wage, rrn?}]
	workplace_info: 사업장 정보(대표자여부 기본값 등 선택) — 없으면 기본값 사용.
	반환: out_path.
	"""
	if not candidates:
		raise ValueError("candidates is empty — 빈 취득신고서를 생성하지 않는다 (fail-closed).")

	workplace_info = workplace_info or {}
	default_representative = str(workplace_info.get("representative_flag", "N"))

	wb = openpyxl.load_workbook(template_path)
	ws = wb[ACQUISITION_SHEET]

	for offset, cand in enumerate(candidates):
		row = ACQUISITION_DATA_START_ROW + offset
		ws[f"A{row}"] = str(cand.get("rrn") or "")
		ws[f"B{row}"] = cand.get("employee_name", "")
		ws[f"C{row}"] = str(cand.get("representative_flag") or default_representative)
		acq_date = _yyyymmdd(cand.get("acquisition_date"))
		wage = int(cand.get("monthly_wage") or 0)
		for col in ACQUISITION_DATE_COLS:
			ws[f"{col}{row}"] = acq_date
		for col in ACQUISITION_WAGE_COLS:
			ws[f"{col}{row}"] = wage

	wb.save(out_path)
	wb.close()
	return out_path


def _resolve_work_days(worker: dict, days_in_month: int) -> list[int]:
	"""근무일 목록 확정. work_days가 있으면 그대로, 없으면 count+last로 역산."""
	work_days = worker.get("work_days")
	if work_days:
		days = sorted(int(d) for d in work_days)
		if days and (days[0] < 1 or days[-1] > days_in_month):
			raise ValueError(f"work_days out of range 1..{days_in_month}: {days}")
		return days
	count = worker.get("work_day_count")
	last = worker.get("last_work_day")
	if count is None or last is None:
		raise ValueError(
			"worker에 work_days 또는 (work_day_count, last_work_day)가 필요하다: "
			f"{worker.get('employee_name')!r}"
		)
	filing = _load_insurance_filing()
	return filing.mark_daily_work_days(int(count), int(last), days_in_month)


def generate_daily_work_report(
	template_path: str,
	workers: list[dict],
	year: int,
	month: int,
	out_path: str,
	workplace_info: dict | None = None,
) -> str:
	"""근로내용확인신고서(일용직) xlsx 생성.

	workers: [{employee_name, rrn?, work_days:[int...] | (work_day_count, last_work_day),
	           daily_wage?, total_wage?, insurance_type?, job_code?, avg_work_hours?,
	           separation_code?}]
	  - work_days가 없고 work_day_count+last_work_day만 있으면 mark_daily_work_days로 역산.
	  - 일자 마킹은 숫자 1 (텍스트 금지 — AQ=COUNT가 텍스트 미집계).
	  - AO/AQ/AS는 수식 유지(값 하드코딩 금지).
	반환: out_path.
	"""
	if not workers:
		raise ValueError("workers is empty — 빈 근로내용확인신고서를 생성하지 않는다 (fail-closed).")

	days_in_month = calendar.monthrange(year, month)[1]
	workplace_info = workplace_info or {}
	default_insurance = str(workplace_info.get("insurance_type", "5"))

	wb = openpyxl.load_workbook(template_path)
	ws = wb[DAILY_SHEET]

	# 템플릿 잔존 데이터 행(PII 포함)을 먼저 비운다.
	last_clear_row = max(DAILY_DATA_LAST_ROW, DAILY_DATA_START_ROW + len(workers) - 1)
	for row in range(DAILY_DATA_START_ROW, last_clear_row + 1):
		for col in range(1, DAILY_CLEAR_LAST_COL + 1):
			ws.cell(row=row, column=col).value = None

	for offset, worker in enumerate(workers):
		row = DAILY_DATA_START_ROW + offset
		work_days = _resolve_work_days(worker, days_in_month)
		ws[f"{DAILY_INSURANCE_COL}{row}"] = str(worker.get("insurance_type") or default_insurance)
		ws[f"{DAILY_NAME_COL}{row}"] = worker.get("employee_name", "")
		ws[f"{DAILY_RRN_COL}{row}"] = str(worker.get("rrn") or "")
		if worker.get("job_code"):
			ws[f"{DAILY_JOB_COL}{row}"] = str(worker["job_code"])
		for day in work_days:
			col = get_column_letter(DAILY_FIRST_DAY_COL_INDEX + day)
			ws[f"{col}{row}"] = 1  # 숫자 1 (COUNT 집계 대상)
		ws[f"{DAILY_WORKDAYS_COL}{row}"] = f"=AQ{row}"
		if worker.get("avg_work_hours") is not None:
			ws[f"{DAILY_AVG_HOURS_COL}{row}"] = worker["avg_work_hours"]
		ws[f"{DAILY_BASE_DAYS_COL}{row}"] = f"=COUNT(J{row}:AN{row})"
		total_wage = worker.get("total_wage")
		if total_wage is None and worker.get("daily_wage") is not None:
			total_wage = int(worker["daily_wage"]) * len(work_days)
		if total_wage is not None:
			ws[f"{DAILY_PAY_TOTAL_COL}{row}"] = int(total_wage)
		ws[f"{DAILY_WAGE_TOTAL_COL}{row}"] = f"=AR{row}"
		if worker.get("separation_code"):
			ws[f"{DAILY_SEPARATION_COL}{row}"] = str(worker["separation_code"])

	wb.save(out_path)
	wb.close()
	return out_path


def generate_loss_report(
	template_path: str,
	candidates: list[dict],
	out_path: str,
	workplace_info: dict | None = None,
) -> str:
	"""상실신고서 xlsx 생성.

	candidates: insurance_filing.detect_losses 출력
	  [{employee, employee_name, loss_date(ISO), loss_reason_code, rrn?}]
	  상실일 = 마지막근무일 + 1 (detect_losses가 이미 계산).
	반환: out_path.
	"""
	if not candidates:
		raise ValueError("candidates is empty — 빈 상실신고서를 생성하지 않는다 (fail-closed).")

	wb = openpyxl.load_workbook(template_path)
	ws = wb[LOSS_SHEET]

	for offset, cand in enumerate(candidates):
		row = LOSS_DATA_START_ROW + offset
		ws[f"A{row}"] = cand.get("employee_name", "")
		ws[f"B{row}"] = str(cand.get("rrn") or "")
		loss_date = _yyyymmdd(cand.get("loss_date"))
		for col in LOSS_DATE_COLS:
			ws[f"{col}{row}"] = loss_date
		ws[f"{LOSS_REASON_COL}{row}"] = str(cand.get("loss_reason_code") or "")

	wb.save(out_path)
	wb.close()
	return out_path
