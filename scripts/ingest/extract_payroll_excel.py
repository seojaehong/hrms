# 기존 급여대장 엑셀 → 시스템 임포트용 JSON 추출 (기존 데이터 임베딩 1/2).
#
# 노무법인 표준 급여대장 포맷(노호 5월분으로 실검증 — 32/32 1원 일치)을 파싱한다.
# 사용:
#   python3 scripts/ingest/extract_payroll_excel.py "<급여대장.xlsx>" 2026-05 [시트명] > payroll.json
# 출력 JSON: {period, employees: [{name, dept, join, email, earnings{}, deductions{},
#            expected_gross, expected_ded, expected_net}]}
# 추출 후 반드시 내부 무결성(earnings 합=세전, 세전-공제=실지급)을 검사하고 불일치면 실패한다.
#
# 주민번호는 추출하지 않는다 (신고서 생성 단계에서만 별도 취급).

from __future__ import annotations

import json
import sys

import openpyxl

# 노무법인 표준 헤더 → 컬럼 인덱스 (1-base). 다른 포맷은 이 매핑만 교체.
EARNING_COLS = {
    20: "기본급", 21: "고정연장수당", 22: "고정야간수당", 23: "고정휴일수당",
    24: "고정휴일연장수당", 25: "연차선지급", 26: "식대(비과세)", 27: "보안수당",
    28: "근로자의날추가지급", 29: "초과근무수당", 30: "전월미지급",
}
DEDUCTION_COLS = {32: "소득세", 33: "지방소득세", 34: "국민연금", 35: "건강보험", 36: "장기요양", 37: "고용보험"}
COL_DEPT, COL_NAME, COL_JOIN, COL_EMAIL = 2, 3, 8, 14
COL_GROSS, COL_DED_TOTAL, COL_NET = 31, 38, 39


def extract(path: str, period: str, sheet: str = "급여(직원)") -> dict:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet]

    def num(r: int, c: int) -> int:
        v = ws.cell(r, c).value
        return int(round(float(v))) if isinstance(v, (int, float)) else 0

    employees = []
    for r in range(2, ws.max_row + 1):
        name = ws.cell(r, COL_NAME).value
        if not name:
            continue
        earnings = {label: num(r, c) for c, label in EARNING_COLS.items() if num(r, c)}
        deductions = {label: num(r, c) for c, label in DEDUCTION_COLS.items() if num(r, c)}
        employees.append(
            {
                "name": str(name).strip(),
                "dept": str(ws.cell(r, COL_DEPT).value or "").strip(),
                "join": str(ws.cell(r, COL_JOIN).value)[:10] if ws.cell(r, COL_JOIN).value else "",
                "email": str(ws.cell(r, COL_EMAIL).value or "").strip(),
                "earnings": earnings,
                "deductions": deductions,
                "expected_gross": num(r, COL_GROSS),
                "expected_ded": num(r, COL_DED_TOTAL),
                "expected_net": num(r, COL_NET),
            }
        )

    errors = []
    for e in employees:
        if sum(e["earnings"].values()) != e["expected_gross"]:
            errors.append(f"{e['name']}: earnings 합 != 세전")
        if e["expected_gross"] - sum(e["deductions"].values()) != e["expected_net"]:
            errors.append(f"{e['name']}: 세전-공제 != 실지급")
    if errors:
        raise SystemExit("무결성 실패 (엑셀이 권위 — 매핑을 점검하라):\n" + "\n".join(errors))

    return {"period": period, "count": len(employees), "employees": employees}


if __name__ == "__main__":
    if sys.platform == "win32":
        sys.stdout.reconfigure(encoding="utf-8")
    if len(sys.argv) < 3:
        raise SystemExit(__doc__)
    data = extract(sys.argv[1], sys.argv[2], sys.argv[3] if len(sys.argv) > 3 else "급여(직원)")
    print(json.dumps(data, ensure_ascii=False, indent=1))
