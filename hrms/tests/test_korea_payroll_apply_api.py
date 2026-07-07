"""급여 엑셀 업로드 확정 API 테스트 — FakeFrappe (frappe 없이 python3 직접 실행).

실행: python3 hrms/tests/test_korea_payroll_apply_api.py
apply_payroll_upload: human_approved 게이트(fail-closed) + Salary Slip upsert 를 검증한다.
  - human_approved != True → status blocked, frappe 무접촉(only_for 호출 없음)·DB 무변경.
  - 승인 경로 → 신규 인원 Slip 생성, 건수·총지급 요약, 감사 Comment 기록.
개인 급여액은 로그하지 않는다(감사 Comment 는 집계 수치만).
"""

from __future__ import annotations

import importlib.util
import io
import pathlib
import sys
import unittest

import openpyxl

MODULE_PATH = pathlib.Path(__file__).resolve().parents[1] / "regional" / "south_korea" / "payroll_excel_api.py"
CORE_PATH = pathlib.Path(__file__).resolve().parents[1] / "regional" / "south_korea" / "payroll_excel.py"


def _load_core():
    spec = importlib.util.spec_from_file_location("korea_payroll_excel_core_apply", CORE_PATH)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


CORE = _load_core()
PROFILE = CORE.PROFILES["monthly"]


def build_workbook(employees: list[dict]) -> bytes:
    """monthly 프로파일 컬럼 매핑으로 인메모리 xlsx 바이트 생성 (업로드 대체)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = PROFILE["sheet"]
    cols = PROFILE["cols"]
    earn_by_label = {label: c for c, label in PROFILE["earnings"].items()}
    ded_by_label = {label: c for c, label in PROFILE["deductions"].items()}
    ws.cell(1, cols["name"]).value = "성명"
    row = 2
    for e in employees:
        ws.cell(row, cols["name"]).value = e["name"]
        for label, amount in e.get("earnings", {}).items():
            ws.cell(row, earn_by_label[label]).value = amount
        for label, amount in e.get("deductions", {}).items():
            ws.cell(row, ded_by_label[label]).value = amount
        ws.cell(row, cols["gross"]).value = e.get("gross", 0)
        ws.cell(row, cols["ded_total"]).value = e.get("ded_total", 0)
        ws.cell(row, cols["net"]).value = e.get("net", 0)
        row += 1
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


class FakeUploadedFile:
    def __init__(self, content: bytes):
        self._content = content

    def get_content(self):
        return self._content


class FakeDoc:
    """get_doc(dict) 반환 — insert 시 tracker 에 기록하고, Employee 는 name 을 부여한다."""

    def __init__(self, data: dict, created: list[dict]):
        self._data = dict(data)
        self._created = created
        self.name = data.get("__newname")

    def insert(self, ignore_permissions=False):
        self._created.append(dict(self._data))
        if self.name is None:
            if self._data.get("doctype") == "Employee":
                self.name = f"EMP-{self._data.get('first_name')}"
            else:
                self.name = self._data.get("salary_component") or "DOC"
        return self

    def submit(self):
        return self

    def __getattr__(self, name):
        try:
            return self._data[name]
        except KeyError as exc:  # pragma: no cover
            raise AttributeError(name) from exc


class FakeDB:
    def __init__(self, default_company, existing_slip_emps=()):
        self._default_company = default_company
        self._existing_slip_emps = set(existing_slip_emps)
        self.commits = 0

    def get_single_value(self, doctype, field):
        if (doctype, field) == ("Global Defaults", "default_company"):
            return self._default_company
        return None

    def exists(self, doctype, filters=None):
        if doctype == "Salary Slip":
            emp = (filters or {}).get("employee")
            return f"SS-{emp}" if emp in self._existing_slip_emps else None
        return None  # 구성항목/구조/배정은 모두 신규로 취급

    def get_value(self, doctype, filters, field):
        if doctype == "Employee":
            return None  # 모든 직원 신규
        return None

    def commit(self):
        self.commits += 1


class FakeFrappe:
    def __init__(self, *, upload_bytes=b"", default_company="노호", allowed_roles=True, existing_slip_emps=()):
        self.db = FakeDB(default_company, existing_slip_emps)
        self._upload_bytes = upload_bytes
        self.allowed_roles = allowed_roles
        self.only_for_calls = []
        self.created_docs = []  # 모든 insert dict — DB 변경 감지용
        self.whitelisted = []

    def whitelist(self):
        def decorator(fn):
            self.whitelisted.append(fn.__name__)
            return fn

        return decorator

    def only_for(self, roles):
        self.only_for_calls.append(tuple(roles))
        if not self.allowed_roles:
            raise PermissionError("not permitted")

    def get_doc(self, first, name=None):
        if isinstance(first, dict):
            return FakeDoc(first, self.created_docs)
        if first == "File":
            return FakeUploadedFile(self._upload_bytes)
        raise AssertionError(f"unexpected get_doc: {first!r}")

    def created_of(self, doctype):
        return [d for d in self.created_docs if d.get("doctype") == doctype]


def load_module(fake_frappe):
    old = sys.modules.get("frappe")
    sys.modules["frappe"] = fake_frappe
    try:
        spec = importlib.util.spec_from_file_location("korea_payroll_apply_api", MODULE_PATH)
        module = importlib.util.module_from_spec(spec)
        assert spec.loader is not None
        spec.loader.exec_module(module)
        return module
    finally:
        if old is not None:
            sys.modules["frappe"] = old
        else:
            sys.modules.pop("frappe", None)


# 자기일관(earnings 합=gross, gross-공제=net) 인 신규 인원 2명
TWO_NEW = [
    {"name": "갑", "earnings": {"기본급": 3000000}, "deductions": {"소득세": 100000},
     "gross": 3000000, "ded_total": 100000, "net": 2900000},
    {"name": "을", "earnings": {"기본급": 2500000, "식대(비과세)": 200000}, "deductions": {"소득세": 80000},
     "gross": 2700000, "ded_total": 80000, "net": 2620000},
]


class TestApplyPayrollUpload(unittest.TestCase):
    def test_blocked_when_not_approved_touches_nothing(self):
        fake = FakeFrappe(upload_bytes=build_workbook(TWO_NEW), default_company="노호")
        module = load_module(fake)
        result = module.apply_payroll_upload(file_url="/private/files/up.xlsx", period="2026-05")
        self.assertEqual(result["status"], "blocked")
        self.assertTrue(result["requires_human_confirmation"])
        # fail-closed: frappe 조회·저장 일절 없음
        self.assertEqual(fake.only_for_calls, [])
        self.assertEqual(fake.created_docs, [])
        self.assertEqual(fake.db.commits, 0)

    def test_blocked_string_false_variants(self):
        fake = FakeFrappe(upload_bytes=build_workbook(TWO_NEW), default_company="노호")
        module = load_module(fake)
        for value in (False, "false", "0", "no", ""):
            result = module.apply_payroll_upload(file_url="/x", period="2026-05", human_approved=value)
            self.assertEqual(result["status"], "blocked", value)
        self.assertEqual(fake.created_docs, [])

    def test_invalid_period_raises_before_gate(self):
        fake = FakeFrappe(default_company="노호")
        module = load_module(fake)
        with self.assertRaisesRegex(ValueError, "YYYY-MM"):
            module.apply_payroll_upload(file_url="/x", period="2026/05", human_approved=True)
        self.assertEqual(fake.only_for_calls, [])
        self.assertEqual(fake.created_docs, [])

    def test_approved_creates_slips_and_summarizes(self):
        fake = FakeFrappe(upload_bytes=build_workbook(TWO_NEW), default_company="노호")
        module = load_module(fake)
        result = module.apply_payroll_upload(
            file_url="/private/files/up.xlsx", period="2026-05", human_approved=True
        )
        self.assertEqual(result["status"], "applied")
        self.assertEqual(result["company"], "노호")
        self.assertEqual(result["created"], 2)
        self.assertEqual(result["skipped"], 0)
        self.assertEqual(result["count"], 2)
        self.assertEqual(result["total_gross"], 3000000 + 2700000)
        # HR Manager 전용 권한 게이트
        self.assertEqual(fake.only_for_calls, [("System Manager", "HR Manager")])
        # Salary Slip 2건 생성 + 감사 Comment 기록 + 커밋
        self.assertEqual(len(fake.created_of("Salary Slip")), 2)
        comments = fake.created_of("Comment")
        self.assertEqual(len(comments), 1)
        self.assertIn("created=2", comments[0]["content"])
        self.assertGreaterEqual(fake.db.commits, 1)

    def test_approved_string_true_coerced(self):
        fake = FakeFrappe(upload_bytes=build_workbook(TWO_NEW), default_company="노호")
        module = load_module(fake)
        result = module.apply_payroll_upload(file_url="/x", period="2026-05", human_approved="true")
        self.assertEqual(result["status"], "applied")
        self.assertEqual(result["created"], 2)

    def test_approved_permission_denied(self):
        fake = FakeFrappe(upload_bytes=build_workbook(TWO_NEW), default_company="노호", allowed_roles=False)
        module = load_module(fake)
        with self.assertRaises(PermissionError):
            module.apply_payroll_upload(file_url="/x", period="2026-05", human_approved=True)
        # 권한 실패 시 어떤 문서도 저장되지 않음
        self.assertEqual(fake.created_docs, [])

    def test_existing_slips_are_skipped(self):
        # 두 직원 모두 이미 해당 월 슬립 보유 → 전부 skipped, Slip 미생성
        fake = FakeFrappe(
            upload_bytes=build_workbook(TWO_NEW),
            default_company="노호",
            existing_slip_emps={"EMP-갑", "EMP-을"},
        )
        module = load_module(fake)
        result = module.apply_payroll_upload(file_url="/x", period="2026-05", human_approved=True)
        self.assertEqual(result["status"], "applied")
        self.assertEqual(result["created"], 0)
        self.assertEqual(result["skipped"], 2)
        self.assertEqual(result["total_gross"], 0)
        self.assertEqual(len(fake.created_of("Salary Slip")), 0)

    def test_integrity_failure_raises(self):
        bad = [{"name": "불일치", "earnings": {"기본급": 3000000}, "deductions": {"소득세": 100000},
                "gross": 3100000, "ded_total": 100000, "net": 3000000}]  # earnings 합 != 세전
        fake = FakeFrappe(upload_bytes=build_workbook(bad), default_company="노호")
        module = load_module(fake)
        with self.assertRaises(ValueError):
            module.apply_payroll_upload(file_url="/x", period="2026-05", human_approved=True)
        self.assertEqual(len(fake.created_of("Salary Slip")), 0)

    def test_no_company_raises(self):
        fake = FakeFrappe(upload_bytes=build_workbook(TWO_NEW), default_company=None)
        module = load_module(fake)
        with self.assertRaisesRegex(ValueError, "company"):
            module.apply_payroll_upload(file_url="/x", period="2026-05", human_approved=True)


class TestSummarizePayrollApply(unittest.TestCase):
    def test_counts_and_total(self):
        results = [
            {"name": "a", "action": "created", "gross": 1000000},
            {"name": "b", "action": "created", "gross": 2000000},
            {"name": "c", "action": "updated", "gross": 500000},
            {"name": "d", "action": "skipped", "gross": 0},
        ]
        summary = CORE.summarize_payroll_apply(results)
        self.assertEqual(summary["created"], 2)
        self.assertEqual(summary["updated"], 1)
        self.assertEqual(summary["skipped"], 1)
        self.assertEqual(summary["count"], 4)
        self.assertEqual(summary["total_gross"], 3500000)

    def test_empty(self):
        summary = CORE.summarize_payroll_apply([])
        self.assertEqual(summary, {"created": 0, "updated": 0, "skipped": 0, "count": 0, "total_gross": 0})


if __name__ == "__main__":
    unittest.main(verbosity=2)
