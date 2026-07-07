"""급여 엑셀 업로드 검증 API 테스트 — FakeFrappe (frappe 없이 python3 직접 실행).

실행: python3 hrms/tests/test_korea_payroll_upload_api.py
validate_payroll_upload: 권한·File 읽기·파싱·무결성(행번호 오류)·기존 슬립 대비 diff 를 검증한다.
저장이 없어야 하며(File 생성/슬립 upsert 호출 없음), 개인 급여액은 로그하지 않는다.
"""

from __future__ import annotations

import importlib.util
import io
import pathlib
import sys
import unittest

import openpyxl

MODULE_PATH = pathlib.Path(__file__).resolve().parents[1] / "regional" / "south_korea" / "payroll_excel_api.py"
CORE_PATH = pathlib.Path(__file__).resolve().parents[1] / "regional" / "south_korea" / "payroll_excel.py"


def _load_core():
    spec = importlib.util.spec_from_file_location("korea_payroll_excel_core", CORE_PATH)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


CORE = _load_core()
PROFILE = CORE.PROFILES["monthly"]


def build_workbook(employees: list[dict]) -> bytes:
    """monthly 프로파일 컬럼 매핑으로 인메모리 xlsx 바이트 생성 (업로드 대체)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = PROFILE["sheet"]
    cols = PROFILE["cols"]
    earn_by_label = {label: c for c, label in PROFILE["earnings"].items()}
    ded_by_label = {label: c for c, label in PROFILE["deductions"].items()}
    ws.cell(1, cols["name"]).value = "성명"
    row = 2
    for e in employees:
        ws.cell(row, cols["name"]).value = e["name"]
        for label, amount in e.get("earnings", {}).items():
            ws.cell(row, earn_by_label[label]).value = amount
        for label, amount in e.get("deductions", {}).items():
            ws.cell(row, ded_by_label[label]).value = amount
        ws.cell(row, cols["gross"]).value = e.get("gross", 0)
        ws.cell(row, cols["ded_total"]).value = e.get("ded_total", 0)
        ws.cell(row, cols["net"]).value = e.get("net", 0)
        row += 1
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


class FakeChildRow(dict):
    def __getattr__(self, name):
        try:
            return self[name]
        except KeyError as exc:  # pragma: no cover
            raise AttributeError(name) from exc


class FakeSlip:
    def __init__(self, data):
        self._data = dict(data)
        self.earnings = [FakeChildRow(r) for r in data.get("earnings", [])]
        self.deductions = [FakeChildRow(r) for r in data.get("deductions", [])]

    def get(self, key):
        return self._data.get(key)

    def __getattr__(self, name):
        try:
            return self._data[name]
        except KeyError as exc:  # pragma: no cover
            raise AttributeError(name) from exc


class FakeUploadedFile:
    def __init__(self, content: bytes):
        self._content = content

    def get_content(self):
        return self._content


class FakeDB:
    def __init__(self, default_company):
        self._default_company = default_company

    def get_single_value(self, doctype, field):
        if (doctype, field) == ("Global Defaults", "default_company"):
            return self._default_company
        return None


class FakeFrappe:
    def __init__(self, *, upload_bytes=b"", slips=None, default_company="노호", allowed_roles=True):
        self.db = FakeDB(default_company)
        self._upload_bytes = upload_bytes
        self._slips = {s["name"]: s for s in (slips or [])}
        self.allowed_roles = allowed_roles
        self.only_for_calls = []
        self.get_all_calls = []
        self.created_files = []  # 저장 없음 검증용 — 절대 채워지면 안 됨
        self.whitelisted = []

    def whitelist(self):
        def decorator(fn):
            self.whitelisted.append(fn.__name__)
            return fn

        return decorator

    def only_for(self, roles):
        self.only_for_calls.append(tuple(roles))
        if not self.allowed_roles:
            raise PermissionError("not permitted")

    def get_all(self, doctype, filters=None, pluck=None):
        self.get_all_calls.append({"doctype": doctype, "filters": dict(filters or {}), "pluck": pluck})
        if doctype != "Salary Slip":
            raise AssertionError(f"unexpected get_all: {doctype}")
        return list(self._slips.keys())

    def get_doc(self, first, name=None):
        if isinstance(first, dict) and first.get("doctype") == "File":
            self.created_files.append(first)  # 저장 시도 감지
            raise AssertionError("validate_payroll_upload must not create files")
        if first == "File":
            return FakeUploadedFile(self._upload_bytes)
        if first == "Salary Slip":
            return FakeSlip(self._slips[name])
        raise AssertionError(f"unexpected get_doc: {first!r}")


def load_module(fake_frappe):
    old = sys.modules.get("frappe")
    sys.modules["frappe"] = fake_frappe
    try:
        spec = importlib.util.spec_from_file_location("korea_payroll_upload_api", MODULE_PATH)
        module = importlib.util.module_from_spec(spec)
        assert spec.loader is not None
        spec.loader.exec_module(module)
        return module
    finally:
        if old is not None:
            sys.modules["frappe"] = old
        else:
            sys.modules.pop("frappe", None)


# 기존 슬립 2명(동일/변경 대상)
EXISTING_SLIPS = [
    {
        "name": "SS-1", "employee_name": "동일", "department": "관리", "posting_date": "2025-03-01",
        "gross_pay": 3000000, "net_pay": 2900000,
        "earnings": [{"salary_component": "기본급", "amount": 3000000}],
        "deductions": [{"salary_component": "소득세", "amount": 100000}],
    },
    {
        "name": "SS-2", "employee_name": "변경", "department": "매장", "posting_date": "2026-01-10",
        "gross_pay": 2600000, "net_pay": 2530000,
        "earnings": [{"salary_component": "기본급", "amount": 2600000}],
        "deductions": [{"salary_component": "소득세", "amount": 70000}],
    },
]


class TestValidatePayrollUpload(unittest.TestCase):
    def test_ok_returns_diff_without_saving(self):
        upload = build_workbook([
            {"name": "동일", "earnings": {"기본급": 3000000}, "deductions": {"소득세": 100000},
             "gross": 3000000, "ded_total": 100000, "net": 2900000},
            {"name": "변경", "earnings": {"기본급": 2600001}, "deductions": {"소득세": 70000},
             "gross": 2600001, "ded_total": 70000, "net": 2530001},  # 세전·실지급 +1원 (행 무결성 유지)
            {"name": "신규", "earnings": {"기본급": 1500000}, "deductions": {"소득세": 40000},
             "gross": 1500000, "ded_total": 40000, "net": 1460000},
        ])
        fake = FakeFrappe(upload_bytes=upload, slips=EXISTING_SLIPS, default_company="노호")
        module = load_module(fake)

        result = module.validate_payroll_upload(file_url="/private/files/up.xlsx", period="2026-05")

        self.assertEqual(result["status"], "ok")
        self.assertEqual(result["company"], "노호")
        self.assertEqual(result["count"], 3)
        counts = result["diff"]["counts"]
        self.assertEqual(counts, {"new": 1, "changed": 1, "same": 1, "missing": 0})
        self.assertEqual(result["diff"]["changed"][0]["net_delta"], 1)  # 1원 delta
        # 권한 게이트 호출 + 해당 월 슬립 조회 + 저장 없음
        self.assertEqual(fake.only_for_calls, [("System Manager", "HR Manager", "HR User")])
        self.assertEqual(fake.get_all_calls[0]["filters"]["start_date"], "2026-05-01")
        self.assertEqual(fake.created_files, [])

    def test_invalid_period_raises_before_frappe(self):
        fake = FakeFrappe(default_company="노호")
        module = load_module(fake)
        with self.assertRaisesRegex(ValueError, "YYYY-MM"):
            module.validate_payroll_upload(file_url="/x", period="2026/05")
        self.assertEqual(fake.only_for_calls, [])

    def test_permission_denied(self):
        fake = FakeFrappe(default_company="노호", allowed_roles=False)
        module = load_module(fake)
        with self.assertRaises(PermissionError):
            module.validate_payroll_upload(file_url="/x", period="2026-05")

    def test_integrity_failure_returns_structured_errors_with_row_numbers(self):
        upload = build_workbook([
            {"name": "불일치", "earnings": {"기본급": 3000000}, "deductions": {"소득세": 100000},
             "gross": 3100000, "ded_total": 100000, "net": 3000000},  # earnings 합 != 세전
        ])
        fake = FakeFrappe(upload_bytes=upload, default_company="노호")
        module = load_module(fake)
        result = module.validate_payroll_upload(file_url="/x", period="2026-05")
        self.assertEqual(result["status"], "invalid")
        self.assertEqual(result["count"], 1)
        self.assertTrue(result["errors"])
        self.assertIn("[1]", result["errors"][0])  # 행 번호 포함
        self.assertEqual(fake.created_files, [])  # 저장 없음

    def test_parse_error_returns_structured_error(self):
        fake = FakeFrappe(upload_bytes=build_workbook([]), default_company="노호")
        module = load_module(fake)
        result = module.validate_payroll_upload(file_url="/x", period="2026-05", sheet_type="weekly")
        self.assertEqual(result["status"], "parse_error")
        self.assertTrue(result["errors"])
        self.assertIn("weekly", result["errors"][0])

    def test_no_company_treats_all_as_new(self):
        upload = build_workbook([
            {"name": "홀로", "earnings": {"기본급": 1000000}, "deductions": {},
             "gross": 1000000, "ded_total": 0, "net": 1000000},
        ])
        fake = FakeFrappe(upload_bytes=upload, default_company=None)
        module = load_module(fake)
        result = module.validate_payroll_upload(file_url="/x", period="2026-05")
        self.assertEqual(result["status"], "ok")
        self.assertIsNone(result["company"])
        self.assertEqual(result["diff"]["counts"]["new"], 1)
        self.assertEqual(fake.get_all_calls, [])  # 회사 없으면 슬립 조회 생략


if __name__ == "__main__":
    unittest.main(verbosity=2)
