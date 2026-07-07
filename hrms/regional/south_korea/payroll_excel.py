"""급여대장 엑셀 파싱·무결성 검증 코어 — framework-free.

로컬 자산 scripts/ingest/extract_payroll_excel.py 의 검증된 파싱 로직
(노호 5월분 32/32 1원 일치)을 hrms 패키지로 승격해 서버 API가 임포트할 수 있게 한다.
frappe import 없음 — openpyxl 만 사용(서버/로컬 python3 모두 가용).

주민번호는 이 모듈에 들어오지 않는다 — 신고서 생성 단계에서만 별도 취급.
"""

from __future__ import annotations

import io
from typing import Any

import openpyxl

# 시트 종류별 컬럼 매핑 프로파일 (1-base 인덱스). 다른 포맷은 프로파일만 추가.
PROFILES = {
    # 월급제 '급여(직원)' 시트 — 노호 5월 32명 1원 일치.
    "monthly": {
        "sheet": "급여(직원)",
        "cols": {"dept": 2, "name": 3, "join": 8, "email": 14,
                 "gross": 31, "ded_total": 38, "net": 39},
        "earnings": {
            20: "기본급", 21: "고정연장수당", 22: "고정야간수당", 23: "고정휴일수당",
            24: "고정휴일연장수당", 25: "연차선지급", 26: "식대(비과세)", 27: "보안수당",
            28: "근로자의날추가지급", 29: "초과근무수당", 30: "전월미지급",
        },
        "deductions": {32: "소득세", 33: "지방소득세", 34: "국민연금", 35: "건강보험", 36: "장기요양", 37: "고용보험"},
    },
    # 시급제 '파트타임' 시트 — U~X=지급, Y=세전합계, Z~AE=공제, AF=공제합계, AG=차인지급액.
    "parttime": {
        "sheet": "파트타임",
        "cols": {"dept": 2, "name": 3, "join": 8, "email": 14,
                 "gross": 25, "ded_total": 32, "net": 33},
        "earnings": {21: "기본급", 22: "주휴수당", 23: "근로자의날추가지급", 24: "기타수당"},
        "deductions": {26: "소득세", 27: "지방소득세", 28: "국민연금", 29: "건강보험", 30: "장기요양", 31: "고용보험"},
    },
}

SUMMARY_NAMES = {"합계", "총계", "소계", "계"}


def parse_payroll_workbook(source: Any, sheet_type: str = "monthly", sheet: str | None = None) -> list[dict]:
    """급여대장 워크북을 파싱해 직원 행 리스트를 반환(무결성 검사 없음).

    source: 파일 경로(str/Path) 또는 xlsx 바이트(bytes). sheet_type: PROFILES 키.
    반환 각 행: {name, dept, join, email, earnings{}, deductions{},
                 expected_gross, expected_ded, expected_net}
    """
    if sheet_type not in PROFILES:
        raise ValueError(f"알 수 없는 sheet_type: {sheet_type} (가능: {', '.join(PROFILES)})")
    profile = PROFILES[sheet_type]
    cols = profile["cols"]
    earning_cols = profile["earnings"]
    deduction_cols = profile["deductions"]

    handle = io.BytesIO(source) if isinstance(source, (bytes, bytearray)) else source
    wb = openpyxl.load_workbook(handle, data_only=True)
    ws = wb[sheet or profile["sheet"]]

    def num(r: int, c: int) -> int:
        v = ws.cell(r, c).value
        if isinstance(v, bool) or not isinstance(v, (int, float)):
            return 0
        return int(round(float(v)))

    rows = []
    for r in range(2, ws.max_row + 1):
        name = ws.cell(r, cols["name"]).value
        if not name:
            continue
        if str(name).strip() in SUMMARY_NAMES:
            continue  # 합계행은 자기일관이라 무결성 검사를 통과하므로 반드시 이름으로 스킵
        earnings = {label: num(r, c) for c, label in earning_cols.items() if num(r, c)}
        deductions = {label: num(r, c) for c, label in deduction_cols.items() if num(r, c)}
        rows.append(
            {
                "name": str(name).strip(),
                "dept": str(ws.cell(r, cols["dept"]).value or "").strip(),
                "join": str(ws.cell(r, cols["join"]).value)[:10] if ws.cell(r, cols["join"]).value else "",
                "email": str(ws.cell(r, cols["email"]).value or "").strip(),
                "earnings": earnings,
                "deductions": deductions,
                "expected_gross": num(r, cols["gross"]),
                "expected_ded": num(r, cols["ded_total"]),
                "expected_net": num(r, cols["net"]),
            }
        )
    return rows


def validate_payroll_rows(rows: list[dict]) -> list[str]:
    """행별 내부 무결성(earnings 합=세전, 세전-공제=실지급)을 검사해 오류 메시지 리스트를 반환.

    빈 리스트면 통과. 엑셀이 권위 — 불일치는 매핑 점검 신호.
    """
    errors = []
    for idx, e in enumerate(rows, start=1):
        name = e.get("name", f"행{idx}")
        earnings = e.get("earnings") or {}
        deductions = e.get("deductions") or {}
        if not earnings and e.get("expected_gross", 0) == 0 and e.get("expected_net", 0) == 0:
            errors.append(f"[{idx}] {name}: 숫자 셀 없음 — 텍스트 서식 행 의심 (0원 시드 방지)")
            continue
        if sum(earnings.values()) != e.get("expected_gross", 0):
            errors.append(f"[{idx}] {name}: earnings 합 != 세전")
        if e.get("expected_gross", 0) - sum(deductions.values()) != e.get("expected_net", 0):
            errors.append(f"[{idx}] {name}: 세전-공제 != 실지급")
    return errors


def extract_payroll(source: Any, period: str, sheet_type: str = "monthly", sheet: str | None = None) -> dict:
    """파싱 + 무결성 검사. 통과 시 {period, count, employees}, 실패 시 ValueError."""
    rows = parse_payroll_workbook(source, sheet_type, sheet)
    errors = validate_payroll_rows(rows)
    if errors:
        raise ValueError("무결성 실패 (엑셀이 권위 — 매핑을 점검하라):\n" + "\n".join(errors))
    return {"period": period, "count": len(rows), "employees": rows}
