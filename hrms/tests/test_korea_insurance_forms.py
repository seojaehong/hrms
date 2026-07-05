"""4대보험 신고서 xlsx 생성기 테스트 — framework-free.

실행: python3 hrms/tests/test_korea_insurance_forms.py

템플릿 원본은 repo 밖 경로를 참조한다(커밋 금지). 더미 후보(가짜 주민번호)로
생성 후 openpyxl로 재열람해 셀 값을 단언한다. 산출물은 tempfile(자동 삭제).
"""

from __future__ import annotations

import importlib.util
import pathlib
import tempfile
import unittest

import openpyxl

ROOT = pathlib.Path(__file__).resolve().parents[2]
MODULE_PATH = ROOT / "hrms" / "regional" / "south_korea" / "insurance_forms.py"

TEMPLATE_DIR = pathlib.Path(
	"C:/Users/iceam/OneDrive/_10_고객/_active/급여자동화/templates"
)
ACQUISITION_TEMPLATE = TEMPLATE_DIR / "근로자고용취득신고_전자신고용 (3).xlsx"
LOSS_TEMPLATE = TEMPLATE_DIR / "근로자자격상실신고서 (3).xlsx"
DAILY_TEMPLATE = TEMPLATE_DIR / "근로내용확인신고_전자신고용.xlsx"


def load_module():
	spec = importlib.util.spec_from_file_location("korea_insurance_forms", MODULE_PATH)
	module = importlib.util.module_from_spec(spec)
	spec.loader.exec_module(module)
	return module


mod = load_module()

# 더미 후보 2명 — 가짜 주민번호(테스트 전용)
DUMMY_ACQUISITIONS = [
	{
		"employee": "E1",
		"employee_name": "김취득",
		"acquisition_date": "2026-07-15",
		"monthly_wage": 3000000,
		"rrn": "0000009999999",
	},
	{
		"employee": "E2",
		"employee_name": "박입사",
		"acquisition_date": "2026-07-01",
		"monthly_wage": 2500000,
		# rrn 없음 → 빈칸
	},
]


@unittest.skipUnless(ACQUISITION_TEMPLATE.exists(), f"템플릿 없음: {ACQUISITION_TEMPLATE}")
class TestAcquisitionReport(unittest.TestCase):
	def test_generates_file_with_correct_cells(self):
		with tempfile.TemporaryDirectory() as tmp:
			out = str(pathlib.Path(tmp) / "취득_신고서.xlsx")
			returned = mod.generate_acquisition_report(
				str(ACQUISITION_TEMPLATE), DUMMY_ACQUISITIONS, out
			)
			self.assertEqual(returned, out)
			self.assertTrue(pathlib.Path(out).exists())

			wb = openpyxl.load_workbook(out)
			ws = wb["서식"]
			# 데이터 시작 행 3, 후보 2명 → 행 3·4
			self.assertEqual(ws["B3"].value, "김취득")
			self.assertEqual(ws["B4"].value, "박입사")
			# 성명·취득일·보수월액 단언
			self.assertEqual(ws["A3"].value, "0000009999999")
			self.assertIn(ws["A4"].value, (None, ""))  # rrn 없음 → 빈칸
			for col in ("H", "P", "V", "AD"):
				self.assertEqual(ws[f"{col}3"].value, "20260715")
				self.assertEqual(ws[f"{col}4"].value, "20260701")
			for col in ("G", "O", "U", "AC"):
				self.assertEqual(ws[f"{col}3"].value, 3000000)
				self.assertEqual(ws[f"{col}4"].value, 2500000)
			# 대표자여부 기본 'N'
			self.assertEqual(ws["C3"].value, "N")
			wb.close()

	def test_empty_candidates_rejected(self):
		with tempfile.TemporaryDirectory() as tmp:
			out = str(pathlib.Path(tmp) / "빈신고서.xlsx")
			with self.assertRaises(ValueError):
				mod.generate_acquisition_report(str(ACQUISITION_TEMPLATE), [], out)
			self.assertFalse(pathlib.Path(out).exists())


# 더미 상실 후보 2명 — detect_losses 출력 형태(loss_date=마지막근무일+1)
DUMMY_LOSSES = [
	{
		"employee": "E3",
		"employee_name": "이상실",
		"loss_date": "2026-07-16",
		"loss_reason_code": "11",  # 자진퇴사
		"rrn": "0000009999998",
	},
	{
		"employee": "E4",
		"employee_name": "최퇴사",
		"loss_date": "2026-07-01",
		"loss_reason_code": "23",  # 권고사직
		# rrn 없음 → 빈칸
	},
]


@unittest.skipUnless(LOSS_TEMPLATE.exists(), f"템플릿 없음: {LOSS_TEMPLATE}")
class TestLossReport(unittest.TestCase):
	def test_generates_file_with_correct_cells(self):
		with tempfile.TemporaryDirectory() as tmp:
			out = str(pathlib.Path(tmp) / "상실_신고서.xlsx")
			returned = mod.generate_loss_report(
				str(LOSS_TEMPLATE), DUMMY_LOSSES, out
			)
			self.assertEqual(returned, out)
			self.assertTrue(pathlib.Path(out).exists())

			wb = openpyxl.load_workbook(out)
			ws = wb["서식"]
			# 데이터 시작 행 2, 후보 2명 → 행 2·3
			self.assertEqual(ws["A2"].value, "이상실")
			self.assertEqual(ws["A3"].value, "최퇴사")
			self.assertEqual(ws["B2"].value, "0000009999998")
			self.assertIn(ws["B3"].value, (None, ""))  # rrn 없음 → 빈칸
			# 상실일 4보험(F/I/O/T) 동일
			for col in ("F", "I", "O", "T"):
				self.assertEqual(ws[f"{col}2"].value, "20260716")
				self.assertEqual(ws[f"{col}3"].value, "20260701")
			# 상실사유코드 P
			self.assertEqual(ws["P2"].value, "11")
			self.assertEqual(ws["P3"].value, "23")
			wb.close()

	def test_empty_candidates_rejected(self):
		with tempfile.TemporaryDirectory() as tmp:
			out = str(pathlib.Path(tmp) / "빈상실.xlsx")
			with self.assertRaises(ValueError):
				mod.generate_loss_report(str(LOSS_TEMPLATE), [], out)
			self.assertFalse(pathlib.Path(out).exists())


# 더미 일용직 2명 — 1명은 work_days 명시, 1명은 count+last로 역산
from openpyxl.utils import get_column_letter  # noqa: E402

DUMMY_DAILY_WORKERS = [
	{
		"employee_name": "정일용",
		"rrn": "0000009999997",
		"work_days": [3, 5, 10],
		"daily_wage": 100000,
		"job_code": "532",
	},
	{
		"employee_name": "한파트",
		# work_days 없음 → count+last 역산: 마지막 20일에서 4일 → [17,18,19,20]
		"work_day_count": 4,
		"last_work_day": 20,
		"total_wage": 400000,
	},
]


@unittest.skipUnless(DAILY_TEMPLATE.exists(), f"템플릿 없음: {DAILY_TEMPLATE}")
class TestDailyWorkReport(unittest.TestCase):
	def test_generates_file_with_correct_markings(self):
		with tempfile.TemporaryDirectory() as tmp:
			out = str(pathlib.Path(tmp) / "근로내용확인_신고서.xlsx")
			returned = mod.generate_daily_work_report(
				str(DAILY_TEMPLATE), DUMMY_DAILY_WORKERS, 2026, 5, out
			)
			self.assertEqual(returned, out)
			self.assertTrue(pathlib.Path(out).exists())

			wb = openpyxl.load_workbook(out)
			ws = wb["서식"]
			# 데이터 시작 행 2, 2명 → 행 2·3
			self.assertEqual(ws["B2"].value, "정일용")
			self.assertEqual(ws["B3"].value, "한파트")
			self.assertEqual(ws["C2"].value, "0000009999997")
			self.assertIn(ws["C3"].value, (None, ""))  # rrn 없음 → 빈칸

			def marked_days(row):
				days = []
				for d in range(1, 32):
					col = get_column_letter(9 + d)
					if ws[f"{col}{row}"].value == 1:
						days.append(d)
				return days

			# 명시 근무일 그대로
			self.assertEqual(marked_days(2), [3, 5, 10])
			# count+last 역산
			self.assertEqual(marked_days(3), [17, 18, 19, 20])
			# 마킹은 숫자 1(텍스트 아님)
			self.assertIsInstance(ws["L2"].value, int)
			# AO/AQ/AS 수식 유지
			self.assertEqual(ws["AO2"].value, "=AQ2")
			self.assertEqual(ws["AQ2"].value, "=COUNT(J2:AN2)")
			self.assertEqual(ws["AS3"].value, "=AR3")
			# 보수총액: 명시 total 우선 / daily_wage*근무일수 계산
			self.assertEqual(ws["AR2"].value, 300000)  # 100000 * 3
			self.assertEqual(ws["AR3"].value, 400000)
			wb.close()

	def test_residual_pii_rows_cleared(self):
		# 템플릿 잔존 데이터 행(2~8)이 근로자 1명만 써도 모두 비워져야 한다.
		with tempfile.TemporaryDirectory() as tmp:
			out = str(pathlib.Path(tmp) / "pii_clear.xlsx")
			mod.generate_daily_work_report(
				str(DAILY_TEMPLATE), [DUMMY_DAILY_WORKERS[0]], 2026, 5, out
			)
			wb = openpyxl.load_workbook(out)
			ws = wb["서식"]
			for row in range(3, 9):  # 잔존 PII 행
				self.assertIn(ws[f"B{row}"].value, (None, ""))
				self.assertIn(ws[f"C{row}"].value, (None, ""))
			wb.close()

	def test_empty_workers_rejected(self):
		with tempfile.TemporaryDirectory() as tmp:
			out = str(pathlib.Path(tmp) / "빈일용.xlsx")
			with self.assertRaises(ValueError):
				mod.generate_daily_work_report(str(DAILY_TEMPLATE), [], 2026, 5, out)
			self.assertFalse(pathlib.Path(out).exists())




class TestWorkDaysRangeGuard(unittest.TestCase):
	"""리뷰 #7 회귀: 명시 work_days가 월 범위를 벗어나면 거절 (열 오염 방지)."""

	def test_out_of_range_rejected(self):
		import tempfile, pathlib as _pl
		workers = [{"employee_name": "범위밖", "work_days": [35], "daily_wage": 100000}]
		with tempfile.TemporaryDirectory() as tmp:
			with self.assertRaises(ValueError):
				mod.generate_daily_work_report(
					str(DAILY_TEMPLATE), workers, 2026, 7, str(_pl.Path(tmp) / "x.xlsx"), None,
				)


if __name__ == "__main__":
	unittest.main()
